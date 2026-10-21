import os
import imaplib
import email
from email.header import decode_header
import re
from datetime import datetime, timedelta
import openpyxl
import pandas as pd
from supabase import create_client, Client

# -----------------------------------------------------------------------------
# VARIABLES DE ENTORNO Y SANITIZACIÓN DE CREDENCIALES
# -----------------------------------------------------------------------------
GMAIL_USER = os.environ.get("GMAIL_USER", "").strip()
GMAIL_PASS = os.environ.get("GMAIL_PASS", "").strip()

# Limpieza de espacios y barras finales para evitar errores de ruta en Supabase (PGRST125)
SUPABASE_URL = os.environ.get("SUPABASE_URL", "").strip().rstrip("/")
SUPABASE_KEY = os.environ.get("SUPABASE_KEY", "").strip()

# -----------------------------------------------------------------------------
# CONEXIÓN Y BÚSQUEDA EN GMAIL
# -----------------------------------------------------------------------------
def connect_gmail():
    mail = imaplib.IMAP4_SSL("imap.gmail.com")
    mail.login(GMAIL_USER, GMAIL_PASS)
    return mail

def fetch_corte_attachments():
    mail = connect_gmail()
    mail.select("inbox")

    # Búsqueda en los últimos 15 días (leídos o no leídos)
    since_date = (datetime.now() - timedelta(days=15)).strftime("%d-%b-%Y")
    search_criterion = f'(SUBJECT "Corte" SINCE "{since_date}")'
    status, messages = mail.search(None, search_criterion)
    email_ids = messages[0].split()

    if not email_ids:
        print(f"No se encontraron correos con el asunto 'Corte' en los últimos 15 días (desde {since_date}).")
        mail.logout()
        return

    print(f"Encontrados {len(email_ids)} correos en los últimos 15 días. Procesando...")

    for email_id in reversed(email_ids):
        status, data = mail.fetch(email_id, "(RFC822)")

        for response_part in data:
            if isinstance(response_part, tuple):
                msg = email.message_from_bytes(response_part[1])
                
                for part in msg.walk():
                    if part.get_content_maintype() == "multipart" or part.get("Content-Disposition") is None:
                        continue

                    filename = part.get_filename()
                    if filename and (filename.endswith(".xlsx") or filename.endswith(".xls")):
                        print(f"📄 Archivo adjunto encontrado: {filename}")
                        
                        file_data = part.get_payload(decode=True)
                        os.makedirs("downloads", exist_ok=True)
                        filepath = os.path.join("downloads", filename)
                        
                        with open(filepath, "wb") as f:
                            f.write(file_data)
                        
                        process_and_upload(filepath, filename)

    mail.logout()

# -----------------------------------------------------------------------------
# PROCESAMIENTO DEL EXCEL E INSERCIÓN EN SUPABASE
# -----------------------------------------------------------------------------
def process_and_upload(filepath, filename):
    supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)
    fecha_cierre = None

    # 1. Extraer y normalizar la fecha del nombre del archivo (soporta D-M-AAAA y AAAA-M-D)
    match_iso = re.search(r'(\d{4})[-_.](\d{1,2})[-_.](\d{1,2})', filename)
    if match_iso:
        ano, mes, dia = match_iso.groups()
        fecha_cierre = f"{ano}-{mes.zfill(2)}-{dia.zfill(2)}"
    else:
        match_dmy = re.search(r'(\d{1,2})[-_.](\d{1,2})[-_.](\d{4})', filename)
        if match_dmy:
            dia, mes, ano = match_dmy.groups()
            fecha_cierre = f"{ano}-{mes.zfill(2)}-{dia.zfill(2)}"

    if not fecha_cierre:
        print(f"⚠️ No se identificó fecha en el nombre del archivo: {filename}")
        return

    # 2. Verificar si el cierre ya figura registrado en la tabla 'cierres_diarios'
    response = supabase.table("cierres_diarios").select("fecha").eq("fecha", fecha_cierre).execute()
    if response.data:
        print(f"ℹ️ El cierre del {fecha_cierre} ya está registrado. Se omite.")
        return

    print(f"🚀 Leyendo Excel para registrar el cierre del {fecha_cierre}...")

    # 3. Lectura del libro de Excel mediante OpenPyXL (idéntica a la plataforma Streamlit)
    wb = openpyxl.load_workbook(filepath, data_only=True)
    sheet = wb.active
    data = [[sheet.cell(row=r, column=c).value for c in range(1, sheet.max_column + 1)] for r in range(1, sheet.max_row + 1)]
    
    header_idx = next((idx for idx, row in enumerate(data) if 'FECHA' in [str(x).upper() if x else '' for x in row] and 'CONCEPTO' in [str(x).upper() if x else '' for x in row]), None)
    
    if header_idx is None:
        print(f"❌ No se encontró la estructura estándar en el Excel: {filename}")
        return
        
    header = [str(x).strip().upper() if x else '' for x in data[header_idx]]
    col_fecha = header.index('FECHA') if 'FECHA' in header else 1
    col_concepto = header.index('CONCEPTO') if 'CONCEPTO' in header else 2
    col_cantidad = header.index('CANTIDAD') if 'CANTIDAD' in header else 3
    
    col_codigo = 4
    for kw in ['CODIGO', 'CÓDIGO', 'NUMERO', 'NÚMERO']:
        if kw in header:
            col_codigo = header.index(kw)
            break
            
    col_ingreso = header.index('INGRESO') if 'INGRESO' in header else 5
    col_gastos = header.index('GASTOS') if 'GASTOS' in header else 6
    col_saldo = header.index('SALDO') if 'SALDO' in header else 7

    transactions = []
    for row in data[header_idx+1:]:
        if row[col_concepto] is None and row[col_ingreso] is None and row[col_gastos] is None: 
            continue
            
        code = row[col_codigo]
        if code is None or str(code).strip() == '': 
            continue

        concepto = str(row[col_concepto]).strip() if row[col_concepto] else ''
        try: cant = int(row[col_cantidad]) if row[col_cantidad] else 1
        except: cant = 1
            
        ingreso = float(row[col_ingreso]) if row[col_ingreso] else 0.0
        gastos = float(row[col_gastos]) if row[col_gastos] else 0.0
        saldo = float(row[col_saldo]) if row[col_saldo] else 0.0
        
        transactions.append({
            'concepto': concepto, 
            'cantidad': cant, 
            'codigo': int(code) if str(code).isdigit() else 1, 
            'ingreso': ingreso, 
            'gastos': gastos, 
            'saldo': saldo
        })

    # 4. Cálculo de Totales según Código de Pago (1: Efectivo, 2: Daviplata, 3: Nequi, 4: Banco)
    saldo_efectivo = sum(t['ingreso'] - t['gastos'] for t in transactions if t['codigo'] == 1)
    saldo_daviplata = sum(t['ingreso'] - t['gastos'] for t in transactions if t['codigo'] == 2)
    saldo_nequi = sum(t['ingreso'] - t['gastos'] for t in transactions if t['codigo'] == 3)
    saldo_banco = sum(t['ingreso'] - t['gastos'] for t in transactions if t['codigo'] == 4)
    total_caja = saldo_efectivo + saldo_daviplata + saldo_nequi + saldo_banco

    # 5. Insertar Cierre Diario en Supabase
    supabase.table("cierres_diarios").insert({
        "fecha": fecha_cierre,
        "efectivo": saldo_efectivo,
        "nequi": saldo_nequi,
        "daviplata": saldo_daviplata,
        "banco": saldo_banco,
        "total": total_caja,
        "observaciones": f"Cargado automáticamente desde {filename}"
    }).execute()

    # 6. Insertar Transacciones Diarias (Desglose de Ingresos y Gastos)
    transacciones_db = []
    for t in transactions:
        transacciones_db.append({
            "fecha": fecha_cierre,
            "concepto": t['concepto'],
            "cantidad": t['cantidad'],
            "codigo": t['codigo'],
            "ingreso": t['ingreso'],
            "gastos": t['gastos'],
            "saldo": t['saldo']
        })
        
    if transacciones_db:
        supabase.table("transacciones_diarias").insert(transacciones_db).execute()

    print(f"✅ Registrado con éxito el día {fecha_cierre} (Totales de caja y {len(transactions)} transacciones).")

if __name__ == "__main__":
    fetch_corte_attachments()
